"""Re-encode a frozen benchmark suite, then run matched observation ablations.

Reuse proposal and caption inputs, never another checkpoint's 3D predictions.
All GPUs first encode disjoint shards, then serve identical reasoning replicas.
Results are published only after exact sample coverage and final-answer audits.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
import csv
import datetime
import hashlib
import os
from pathlib import Path
import queue
import re
import signal
import subprocess
import sys
import threading
import time

from core.captions import CaptionStore
from core.geometry import GeometryStore
from core.io import digest, read_json, write_json
from core.response_constraints import VERSION, matches_constraint
from core.runner import output_lock
from scripts.finish_scannet_recovery import process_identity
from scripts.gpu_lease import acquire_gpu
from scripts.reasoner_pool import ReasonerPool, wait_for_free_gpu

ROOT = Path(__file__).resolve().parents[1]
MERGED = 'merged_bbox_3dconf_top5meanwhl1r_timestamp'


def merge_geometry(paths, expected):
    scenes = {}
    for path in paths:
        items = GeometryStore(path).scenes
        if scenes.keys() & items.keys():
            raise ValueError('Duplicate geometry sample IDs across shards')
        scenes.update(items)
    if set(scenes) != set(expected):
        raise ValueError('Incomplete or unexpected geometry sample IDs')
    return {'scenes': scenes}


def audit_result(result, expected, mode, caption_digest=None):
    rows = result['results']; config = result['config']
    if (len(rows) != len(expected) or {r['sample_id'] for r in rows} != set(expected)
            or result['num_samples'] != len(expected)
            or config.get('answer_format') != VERSION
            or config['observations']['mode'] != mode):
        raise ValueError('QA sample coverage or answer protocol mismatch')
    if caption_digest is not None and config['observations']['captions']['digest'] != caption_digest:
        raise ValueError('Frozen caption bundle changed')
    if any(r['status'] != 'ok' or r.get('finish_reason') not in {'stop', 'eos_token'}
           or not r.get('response_constraint')
           or not matches_constraint(r['raw_output'], r['response_constraint']) for r in rows):
        raise ValueError('Invalid or truncated final answers')


def scannet_metrics(path):
    text = Path(path).read_text()
    if 'Evaluation Results:' not in text or '--- Macro F1' not in text:
        raise ValueError('Incomplete ScanNet metrics log')
    values = {k: float(v) for k, v in re.findall(
        r'^\s+([\w@.]+):\s+([0-9.]+)\s*$', text.split('Evaluation Results:')[-1], re.M)}
    macro = text.split('--- Macro F1')[1].split('Evaluation Results:')[0]
    values.update({f'macro_f1@{k}': float(v) for k, v in re.findall(r'F1@([\d.]+):\s+([\d.]+)', macro)})
    categories = {}
    for line in text.splitlines():
        parts = [p.strip() for p in line.split('|')]
        if len(parts) == 5 and re.fullmatch(r'\d+\s+\d+\s+\d+\s+\d+', parts[1]):
            categories[parts[0]] = {f'F1@{threshold}': float(part.split()[-1])
                                   for threshold, part in zip((.1, .25, .5), parts[2:])}
    return values, categories


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', required=True)
    parser.add_argument('--output-root', required=True)
    parser.add_argument('--validate-only', action='store_true', help='Validate frozen inputs without launching GPU jobs')
    args = parser.parse_args()
    c = read_json(args.config); out = Path(args.output_root).resolve()
    out.mkdir(parents=True, exist_ok=True)
    repo = Path(c['model_repo']); gpus = c['gpus']; shards = len(gpus)
    if not gpus or len(set(gpus)) != shards:
        raise ValueError('GPU IDs must be nonempty and unique')
    sizes = [item for key in ('resolution', 'spatial_resolution', 'model_image_size')
             for item in ('--' + key.replace('_', '-'), str(c[key]))]
    state = {'pid': os.getpid(), 'phase': 'preflight', 'encoding': {}, 'tasks': {}, 'workers': {}}
    guard = threading.RLock(); children = []; pool = None

    def save(**fields):
        with guard:
            state.update(fields, updated=time.time()); write_json(out/'status.json', state)

    def update(group, key, **fields):
        with guard:
            state[group].setdefault(str(key), {}).update(fields); save()

    def run(command, log, gpu=None, cwd=ROOT, extra_env=None):
        env = dict(os.environ, OMP_NUM_THREADS='2', MKL_NUM_THREADS='2', OPENBLAS_NUM_THREADS='2',
                   PYTHONUNBUFFERED='1', PYTORCH_CUDA_ALLOC_CONF='expandable_segments:True',
                   PYTHONPATH=str(ROOT), TOKENIZERS_PARALLELISM='false')
        if gpu is not None:
            env['CUDA_VISIBLE_DEVICES'] = str(gpu)
        env.update(extra_env or {}); log = Path(log); log.parent.mkdir(parents=True, exist_ok=True)
        with log.open('ab') as stream:
            p = subprocess.Popen(list(map(str, command)), cwd=cwd, env=env, stdin=subprocess.DEVNULL,
                                 stdout=stream, stderr=subprocess.STDOUT, start_new_session=True)
        with guard:
            children.append(p)
        if p.wait() != 0:
            raise RuntimeError(f'Command failed; inspect {log}')

    def stop_children():
        for child in children:
            if child.poll() is None:
                try: os.killpg(child.pid, signal.SIGTERM)
                except ProcessLookupError: pass
        deadline = time.monotonic() + 30
        for child in children:
            if child.poll() is None:
                try: child.wait(timeout=max(.1, deadline-time.monotonic()))
                except subprocess.TimeoutExpired:
                    os.killpg(child.pid, signal.SIGKILL); child.wait()

    def interrupted(signum, frame):
        save(phase='interrupted', signal=signum)
        stop_children()
        if pool: pool.stop()
        os._exit(128+signum)

    signal.signal(signal.SIGINT, interrupted); signal.signal(signal.SIGTERM, interrupted)

    def publish_summary():
        rows = [{'benchmark': key.split('/')[0], 'variant': c['variant'],
                 'mode': key.split('/')[-1], **value} for key, value in state['tasks'].items()]
        write_json(out/'summary.json', rows)
        columns = ['benchmark', 'variant', 'mode', 'phase', 'samples', 'score_100', 'result', 'error']
        with (out/'summary.csv').open('w') as stream:
            writer = csv.DictWriter(stream, columns, extrasaction='ignore'); writer.writeheader(); writer.writerows(rows)
        if c.get('baseline_table'):
            with open(c['baseline_table'], encoding='utf-8-sig') as stream:
                previous = list(csv.DictReader(stream))
            names = {j['table_name']: j['name'] for j in c['jobs']}
            if c.get('scannet'): names[c['scannet']['table_name']] = c['scannet']['name']
            new_cols = {'boxes': 'WDS518 e100 boxes', 'rgb_boxes': 'RGB + WDS518 e100',
                        'caption_boxes': 'Caption + WDS518 e100'}
            metrics_path = out/'scannet/metrics.log'
            if c.get('scannet') and state['encoding'].get(c['scannet']['name'], {}).get('phase') == 'complete':
                values, categories = scannet_metrics(metrics_path)
                for row in previous:
                    value = None
                    if row['Section'] == 'Direct 3D geometry':
                        value = values.get(row['Metric'])
                        if value is not None and '@' in row['Metric']: value *= 100
                    elif row['Section'] == '3D geometry category':
                        category = row['Benchmark / split'].split(' / ', 1)[-1]
                        value = categories.get(category, {}).get(row['Metric'])
                        if value is not None: value *= 100
                    if value is not None:
                        row[new_cols['boxes']] = value
                        row['Sources'] += '; ' + str(metrics_path)
            for row in previous:
                if row['Section'] != 'Current benchmark': continue
                name = names.get(row['Benchmark / split'])
                if not name: continue
                phases = []
                for mode, col in new_cols.items():
                    task = state['tasks'].get(f'{name}/{c["variant"]}/{mode}', {})
                    row[col] = task.get('score_100') if task.get('phase') == 'complete' else ''
                    if task: phases.append(mode + '=' + task['phase'])
                    if task.get('result'): row['Sources'] += '; ' + task['result']
                row['WDS518 evaluation status'] = '; '.join(phases)
                row['Protocol / caveats'] += ' New Small: WDS518 epoch100, encoder RGB/spatial/internal size 518. Old encoders used 1024/504/1008; this compares checkpoint configurations, not weights alone.'
            fields = list(previous[0]) + [v for v in new_cols.values() if v not in previous[0]]
            if 'WDS518 evaluation status' not in fields: fields.append('WDS518 evaluation status')
            with (out/'comparison_all_results.csv').open('w', encoding='utf-8-sig') as stream:
                writer = csv.DictWriter(stream, fields); writer.writeheader(); writer.writerows(previous)
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
            except ImportError:
                pass  # CSV/Markdown remain available in minimal installations.
            else:
                book = Workbook(); sheet = book.active; sheet.title = 'All results'
                sheet.append(fields)
                for row in previous: sheet.append([row.get(k) for k in fields])
                sheet.freeze_panes = 'I2'; sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='24476B')
                for col in sheet.columns:
                    sheet.column_dimensions[col[0].column_letter].width = min(40, max(16, len(str(col[0].value))+2))
                temporary = out/'comparison_all_results.tmp.xlsx'; book.save(temporary)
                temporary.replace(out/'comparison_all_results.xlsx')
            display = ['Benchmark / split', 'RGB', 'Full boxes', 'Small boxes', *new_cols.values(),
                       'RynnBrain1.1 9B (reported)', 'PhysBrain1.5 8B (reported)']
            def fmt(value):
                if value in (None, ''): return '—'
                try: return f'{float(value):.2f}'
                except (ValueError, TypeError): return str(value).replace('|', '/')
            lines = ['# New checkpoint comparison', '',
                     'Only complete, audited new arms receive scores. Blank cells are pending or unavailable. '
                     'The CSV preserves all previous result rows. Paper scores are references with different protocols.', '',
                     '| ' + ' | '.join(display) + ' |', '| ' + ' | '.join(['---']*len(display)) + ' |']
            lines += ['| ' + ' | '.join(fmt(row.get(k)) for k in display) + ' |'
                      for row in previous if row['Section'] == 'Current benchmark']
            (out/'comparison.md').write_text('\n'.join(lines) + '\n')

    with output_lock(out/'status.json'):
        if (out/'config.json').exists() and read_json(out/'config.json') != c:
            raise ValueError('Configuration changed; use a new output directory')
        write_json(out/'config.json', c)
        write_json(out/'process.json', {'pid': os.getpid(), 'start_time': process_identity(os.getpid()),
            'started_utc': datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'command': sys.argv, 'launcher_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest()})
        try:
            expected = {}; inputs = {}; tasks = queue.Queue()
            for job in c['jobs']:
                name = job['name']; manifest = read_json(job['manifest'])
                ids = [r['sample_id'] for r in manifest['samples']]
                if len(set(ids)) != len(ids): raise ValueError('Duplicate manifest IDs')
                expected[name] = set(ids); index = read_json(Path(job['proposals'])/'index.json')
                if set(index['config']['sample_ids']) != set(ids) or len(index['records']) != len(ids):
                    raise ValueError(f'Frozen proposal coverage mismatch: {name}')
                inputs[name] = {'manifest': digest(manifest), 'proposal_index': digest(index)}
                if job.get('captions'):
                    CaptionStore(job['captions'], manifest)
                    inputs[name]['captions'] = digest(read_json(job['captions']))
                for mode in ('boxes', 'rgb_boxes', 'caption_boxes'):
                    update('tasks', f'{name}/{c["variant"]}/{mode}',
                           phase='queued' if mode != 'caption_boxes' or job.get('captions') else 'unavailable',
                           **({'error': 'No complete frozen caption bundle'} if mode == 'caption_boxes' and not job.get('captions') else {}))
                for shard in range(shards): tasks.put((job, shard))
            scan = c.get('scannet')
            if scan:
                expected[scan['name']] = {r['sample_id'] for r in read_json(scan['manifest'])['samples']}
                inputs[scan['name']] = {'manifest': digest(read_json(scan['manifest']))}
                update('tasks', f'{scan["name"]}/{c["variant"]}/boxes', phase='queued')
            if (out/'input_digests.json').exists() and read_json(out/'input_digests.json') != inputs:
                raise ValueError('Frozen inputs changed since the previous run')
            write_json(out/'input_digests.json', inputs); publish_summary()
            if args.validate_only:
                save(phase='validated', media_samples=sum(len(expected[j['name']]) for j in c['jobs']),
                     planned_arms=sum(v['phase'] == 'queued' for v in state['tasks'].values()))
                print('Frozen input validation passed', flush=True)
                return

            def encode_worker(gpu):
                lease = acquire_gpu(gpu)
                try:
                    wait_for_free_gpu(gpu, lambda g, **kw: update('workers', g, **kw))
                    while True:
                        try: job, shard = tasks.get_nowait()
                        except queue.Empty: break
                        name = job['name']; key = f'{name}/shard_{shard}'
                        dest = out/'media'/name/f'shard_{shard}'
                        update('encoding', key, phase='encoding', gpu=gpu, started=time.time())
                        try:
                            run([c['encoder_python'], '-u', '-m', 'scripts.encode_media_geometry',
                                 '--manifest', job['manifest'], '--proposals', job['proposals'],
                                 '--model-repo', repo, '--checkpoint', c['checkpoint'], '--profile', c['profile'],
                                 '--output', dest, '--shard-index', shard, '--num-shards', shards, *sizes],
                                out/'logs'/f'{name}_shard_{shard}.log', gpu)
                            update('encoding', key, phase='complete', completed=time.time())
                        except Exception as exc:
                            update('encoding', key, phase='failed', error=str(exc))
                        finally: tasks.task_done()
                finally: lease.close()

            save(phase='encoding_media')
            with ThreadPoolExecutor(shards) as workers: list(workers.map(encode_worker, gpus))
            for job in c['jobs']:
                name = job['name']
                try:
                    if any(state['encoding'][f'{name}/shard_{s}']['phase'] != 'complete' for s in range(shards)):
                        raise ValueError('At least one geometry shard failed')
                    bundle = merge_geometry([out/'media'/name/f'shard_{s}'/'geometry.json'
                                             for s in range(shards)], expected[name])
                    write_json(out/'media'/name/'geometry.json', bundle)
                    update('encoding', name, phase='complete', samples=len(bundle['scenes']))
                except Exception as exc:
                    update('encoding', name, phase='failed', error=str(exc))

            if scan:
                save(phase='encoding_scannet')
                scan_out = out/'scannet'; save_name = c['variant'] + '_scannet_val88_gt2d_all'
                scene_list = repo/'inference_gt2d/scannet_val88.txt'; scenes = scene_list.read_text().split()
                scan_env = {'BOX_DATA_PATH': scan['data_root'], 'BOX_OUTPUT_PATH': str(scan_out),
                            'FOURDEVAL_TRACKER_MASK_OFFLOAD': '1', 'FOURDEVAL_LARGE_INTERPOLATION': '1',
                            'FOURDEVAL_CHUNKED_ROPE': '1', 'FOURDEVAL_INPLACE_ROPE': '1'}
                try:
                    run([sys.executable, repo/'inference_gt2d/make_balanced_shards.py',
                         '--scene-list', scene_list, '--data-root', scan['data_root'],
                         '--output-dir', scan_out/'shards', '--shards', shards], scan_out/'shards.log')
                    def encode_scan(item):
                        shard, gpu = item; lease = acquire_gpu(gpu)
                        try:
                            wait_for_free_gpu(gpu, lambda g, **kw: update('workers', g, **kw))
                            run([c['encoder_python'], ROOT/'scripts/encoder_entry.py', repo/'inference_gt2d/scene_inference.py',
                                 scan_out/'shards'/f'shard_{shard:02d}.txt', save_name, 'scannet',
                                 '--model-profile', c['profile'], '--checkpoint', c['checkpoint'],
                                 '--device', 'cuda', '--num-workers', '2', '--max-objects', '1000',
                                 '--object-filter-mode', 'gt2d_all', '--skip-data-validation', '--skip-done',
                                 '--data-root', scan['data_root'], '--output-root', scan_out, *sizes],
                                scan_out/'logs'/f'encoder_{shard}.log', gpu, repo, scan_env)
                        finally: lease.close()
                    with ThreadPoolExecutor(shards) as workers: list(workers.map(encode_scan, enumerate(gpus)))
                    pred = scan_out/'BoxDet/pred'/save_name
                    if {p.parent.name for p in pred.glob('*/done.flag')} != set(scenes):
                        raise ValueError('Not all ScanNet scenes finished')
                    save(phase='scannet_metrics')
                    def merge_scan(scene):
                        run([c['encoder_python'], repo/'inference_gt2d/merge_json.py', '--scene_id', scene,
                             '--pred_dir', save_name, '--exp_name', MERGED], scan_out/'logs'/f'merge_{scene}.log',
                            cwd=repo, extra_env=scan_env)
                    with ThreadPoolExecutor(shards) as workers: list(workers.map(merge_scan, scenes))
                    run([c['encoder_python'], repo/'inference_gt2d/compute_metrics_3d.py', 'scannet',
                         '--pred-dir', save_name, '--exp-name', MERGED], scan_out/'metrics.log', cwd=repo, extra_env=scan_env)
                    run([sys.executable, '-m', 'scripts.convert_geometry', '--format', 'merged',
                         '--input', pred, '--output', scan_out/'geometry.json', '--merged-name', MERGED+'.json',
                         '--quaternion-order', 'xyzw', '--units', 'm', '--coordinate-frame', 'scannet_world_z_up',
                         '--geometry-source', 'predicted', '--encoder', 'spatial_encoder_v2_small',
                         '--checkpoint', c['checkpoint'], '--proposal-source', 'gt_2d',
                         '--label-source', 'gt_category', '--camera-pose-source', 'ground_truth'], scan_out/'convert.log')
                    if set(GeometryStore(scan_out/'geometry.json').scenes) != set(scenes):
                        raise ValueError('ScanNet geometry does not cover all scenes')
                    update('encoding', scan['name'], phase='complete', scenes=len(scenes))
                    publish_summary()
                except Exception as exc:
                    update('encoding', scan['name'], phase='failed', error=str(exc))

            pool = ReasonerPool(c, out, lambda gpu, **kw: update('workers', gpu, **kw))
            save(phase='starting_reasoners'); endpoints = pool.start(); save(phase='reasoning')
            for job in c['jobs'] + ([scan] if scan else []):
                name = job['name']; is_scan = job is scan
                for mode in (['boxes'] if is_scan else ['boxes', 'rgb_boxes', 'caption_boxes']):
                    key = f'{name}/{c["variant"]}/{mode}'
                    if state['tasks'][key]['phase'] == 'unavailable': continue
                    if state['encoding'].get(name, {}).get('phase') != 'complete':
                        update('tasks', key, phase='failed', error='Geometry incomplete'); publish_summary(); continue
                    target = out/'qa'/key/'qa.json'
                    geometry = out/'scannet/geometry.json' if is_scan else out/'media'/name/'geometry.json'
                    try:
                        update('tasks', key, phase='running', started=time.time())
                        command = [sys.executable, 'eval.py', '--benchmark', job['benchmark'], '--data', job['data'],
                            '--split', job['split'], '--model', c['llm_name'], '--base-urls', *endpoints,
                            '--temperature', '0', '--seed', '0', '--max-tokens', '4096', '--timeout', '300',
                            '--concurrency', str(c.get('qa_concurrency', 32)), '--batch-size', '64',
                            '--answer-format', 'native', '--observation-mode', mode, '--media-manifest', job['manifest'],
                            '--geometry', geometry, '--geometry-decimals', '4', '--max-prompt-chars', '600000',
                            '--max-image-pixels', str(c.get('max_image_pixels', 262144)), '--video-frames', '16',
                            '--extra-body', '{"chat_template_kwargs":{"enable_thinking":false}}',
                            '--run-label', key, '--output', target, '--resume', '--verbose']
                        if job.get('dataset'): command += ['--dataset', job['dataset']]
                        if mode == 'caption_boxes': command += ['--captions', job['captions']]
                        run(command, target.parent/'eval.log')
                        result = read_json(target)
                        audit_result(result, expected[name], mode,
                                     inputs[name]['captions'] if mode == 'caption_boxes' else None)
                        update('tasks', key, phase='complete', samples=result['num_samples'],
                               score_100=result['primary_metric']['score_100'], result=str(target), completed=time.time())
                    except Exception as exc: update('tasks', key, phase='failed', error=str(exc))
                    publish_summary()
            failed = [k for k, v in state['tasks'].items() if v['phase'] not in {'complete', 'unavailable'}]
            unavailable = [k for k, v in state['tasks'].items() if v['phase'] == 'unavailable']
            save(phase='failed' if failed else ('complete_with_unavailable_arms' if unavailable else 'complete'),
                 failed_arms=failed, unavailable_arms=unavailable)
            if failed: raise RuntimeError('Some evaluation arms failed; inspect status and logs')
        except Exception as exc:
            save(phase='failed', error=f'{type(exc).__name__}: {exc}'); raise
        finally:
            stop_children()
            if pool: pool.stop()


if __name__ == '__main__':
    main()
