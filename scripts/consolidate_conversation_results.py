"""Consolidate the conversation's evaluation results into one auditable table.

Exports a single-sheet workbook, CSV, JSON and Markdown. Historical protocols,
geometry metrics and published references have explicit row types and units.
"""
import argparse
import collections
import csv
import datetime
import math
from pathlib import Path
import re
import statistics
from zoneinfo import ZoneInfo

from core.io import read_json, write_json

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT.parents[1] / 'physbrain1.5_tech_report.pdf'
BENCHES = [
    ('q_spatial_bench_QSpatial_plus', 'Q-Spatial-plus', 'static image', 'Q-Spatial-Bench'),
    ('blink_val', 'BLINK spatial subset', 'static images', 'BLINK'),
    ('cv_bench_test', 'CV-Bench', 'static image', 'CV-Bench'),
    ('embspatial_bench_test', 'EmbSpatial-Bench', 'static image', 'EmbSpatial-Bench'),
    ('mindcube_tinybench', 'MindCube tinybench', 'multiple views', 'MindCube'),
    ('mmsi_bench_test', 'MMSI-Bench', 'multiple images', 'MMSI-Bench'),
    ('viewspatial_bench_test', 'ViewSpatial-Bench', 'multiple views', 'ViewSpatial-Bench'),
    ('3dsrbench_test', '3DSRBench', 'static image / transformations', '3DSRBench'),
    ('sat_test', 'SAT', 'spatial / multiple views', 'SAT'),
    ('sti_bench_train', 'STI-Bench train', 'dynamic video', None),
    ('vlm4d_real_mc', 'VLM4D real MC', 'dynamic video', None),
    ('vlm4d_synthetic_mc', 'VLM4D synthetic MC', 'dynamic video', None),
    ('dsi_bench_all', 'DSI-Bench all augmentations', 'object / camera dynamics', None),
]
OUR_COLUMNS = ['RGB', 'Full boxes', 'Small boxes', 'RGB + Full', 'RGB + Small',
               'Caption', 'Caption + Full', 'Caption + Small', 'GT boxes']
PAPER_MODELS = ['Gemini 3.6 Flash (reported)', 'GPT 6 Astra (reported)',
                'Claude Opus 5 (reported)', 'Hy-Emb.-VLM-1.0 30A3B (reported)',
                'Embodied-R1.5 8B (reported)', 'RynnBrain1.1 9B (reported)',
                'RoboBrain2.5 8B (reported)', 'MiMo Embodied 7B (reported)',
                'Cosmos3 Nano 8B+8B (reported)', 'ACE-Brain-0.5 8B (reported)',
                'Qwen3-VL-Instruct 8B (reported)', 'PhysBrain1.5 8B (reported)']
PAPER_BENCHES = ['BLINK', 'CV-Bench', '3DSRBench', 'EmbSpatial-Bench', 'MindCube',
    'MMSI-Bench', 'Q-Spatial-Bench', 'RoboSpatial-Home', 'SAT', 'VSI-Bench',
    'ViewSpatial-Bench', 'COSMOS', 'EgoPlan-Bench2', 'ERQA', 'ERQA-PLUS', 'RoboVQA',
    'VLABench', 'Part-Affordance', 'PIOBench', 'PixMo-Points', 'PointBench',
    'RefSpatial-Bench', 'RoboAfford', 'RoboRefit', 'VABench-Point', 'Where2Place',
    'ShareRobot-Traj.', 'VABench-V.-Trace', 'Overall Average']
ARM = {('shared_rgb', 'rgb'): 'RGB', ('full_epoch64', 'boxes'): 'Full boxes',
       ('small_epoch34', 'boxes'): 'Small boxes', ('full_epoch64', 'rgb_boxes'): 'RGB + Full',
       ('small_epoch34', 'rgb_boxes'): 'RGB + Small', ('shared_caption', 'caption'): 'Caption',
       ('full_epoch64', 'caption_boxes'): 'Caption + Full',
       ('small_epoch34', 'caption_boxes'): 'Caption + Small', ('gt3d_oracle', 'boxes'): 'GT boxes'}
META_COLUMNS = ['Section', 'Benchmark / split', 'Scene / input', 'Metric', 'Units',
                'Direction', 'N / unit', 'Status']
COLUMNS = META_COLUMNS + OUR_COLUMNS + PAPER_MODELS + ['Protocol / caveats', 'Sources']


def paper_values():
    import fitz
    with fitz.open(PAPER) as doc:
        lines = doc[10].get_text(sort=True).splitlines()
    result = {}
    for benchmark in PAPER_BENCHES:
        hits = []
        for line in lines:
            match = re.search(r'(?<![\w-])' + re.escape(benchmark) + r'\s+', line)
            if match:
                values = re.findall(r'(?<![\w.])\d+\.\d+(?![\w.])', line[match.end():])
                if len(values) == 12:
                    hits.append([float(v) for v in values])
        assert len(hits) == 1, (benchmark, hits)
        result[benchmark] = dict(zip(PAPER_MODELS, hits[0]))
    assert result['BLINK']['RynnBrain1.1 9B (reported)'] == 84.8
    assert result['VSI-Bench']['PhysBrain1.5 8B (reported)'] == 61.9
    return result


def make_row(section, benchmark, scene, metric, n='', status='complete',
             units='0–100', direction='higher', notes='', sources=()):
    return {'Section': section, 'Benchmark / split': benchmark, 'Scene / input': scene,
            'Metric': metric, 'Units': units, 'Direction': direction, 'N / unit': n,
            'Status': status, 'Protocol / caveats': notes,
            'Sources': '; '.join(str(p) for p in sources)}


def safe_md(value):
    if value is None or value == '':
        return '—'
    if isinstance(value, float):
        return f'{value:.2f}'
    return str(value).replace('|', '/').replace('\n', ' ')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'results/conversation_results')
    args = parser.parse_args(); out = args.output_dir.resolve(); out.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.datetime.now(ZoneInfo('America/Los_Angeles')).isoformat()
    paper = paper_values()
    main_root = ROOT / 'results/observation_ablations_20260915'
    summary = read_json(main_root / 'summary.json')
    final_audit = read_json(main_root / 'final_answer_audit.json')
    assert final_audit['passed'] and final_audit['tasks'] == 68 and final_audit['responses'] == 192813
    caption_roots = [ROOT / 'results/caption_ablations_20260916',
                     ROOT / 'results/caption_recovery_20260916/rerun',
                     ROOT / 'results/caption_recovery_20261001/rerun']
    captions = {}
    for source in caption_roots:
        if not (source / 'summary.json').exists():
            continue
        grouped = collections.defaultdict(list)
        for record in read_json(source / 'summary.json'):
            grouped[record['benchmark']].append(record)
        for benchmark, records in grouped.items():
            if len(records) == 3 and all(r['phase'] == 'complete' for r in records):
                assert {(r['variant'], r['mode']) for r in records} == {
                    ('shared_caption', 'caption'), ('full_epoch64', 'caption_boxes'),
                    ('small_epoch34', 'caption_boxes')}
                captions[benchmark] = records
    selected = summary + [r for group in captions.values() for r in group]
    assert len({(r['benchmark'], r['variant'], r['mode']) for r in selected}) == len(selected)
    groups = collections.defaultdict(list)
    for record in selected:
        groups[record['benchmark']].append(record)
    cache = {}
    def result(record):
        path = record['result']
        if path not in cache:
            d = read_json(path)
            assert len(d['results']) == d['num_samples'] == record['samples']
            assert math.isclose(d['primary_metric']['score_100'], record['score_100'], abs_tol=1e-8)
            assert all(r['status'] == 'ok' and r['finish_reason'] in {'stop', 'eos_token'} for r in d['results'])
            cache[path] = {'statistics': d['statistics'], 'config': d['config'],
                           'question_type_counts': dict(collections.Counter(
                               r.get('question_type') for r in d['results']))}
        return cache[path]
    def fill(row, records):
        for record in records:
            result(record)
            row[ARM[(record['variant'], record['mode'])]] = record['score_100']
        return row
    rows = []; primary = []
    shared_notes = ('Our reasoner: Qwen3.5-9B, temperature 0, seed 0; native final answers, '
                    'max 4096 tokens. Reported paper scores use different visual/prompt/split settings; '
                    'reference only, not a controlled ranking. RGB videos use at most 16 frames.')
    for name, label, scene, paper_name in BENCHES:
        records = groups[name]; n = records[0]['samples']; assert all(r['samples'] == n for r in records)
        note = shared_notes
        if name == '3dsrbench_test':note += ' 11,686 inference rows; grouped circular/flip scoring on 5,157 groups.'
        if name == 'dsi_bench_all':note += ' 7,076 augmented rows from 1,769 base items.'
        if name in captions:
            cap = next(r for r in records if r['mode'] == 'caption')
            settings = result(cap)['config']['observations']['captions']
            penalty = settings['generator']['extra_body'].get('repetition_penalty', 1.0)
            digests = {result(r)['config']['observations']['captions']['digest'] for r in captions[name]}
            assert len(digests) == 1
            note += f' Shared question-independent captions; generation repetition penalty {penalty}.'
        else:
            note += ' Caption arms have no complete score; blank values are not zero.'
            latest = read_json(caption_roots[-1] / 'status.json')['jobs'].get(name, {})
            note += f" Latest caption state: {latest.get('phase')}; valid={latest.get('caption_complete', 0)}, failed={latest.get('caption_failed', 0)}."
        row = fill(make_row('Current benchmark', label, scene, 'native benchmark score',
            f'{n} questions / arm', 'complete' if name in captions else 'caption incomplete',
            notes=note, sources=[r['result'] for r in records] + ([f'{PAPER}#page=11'] if paper_name else [])), records)
        if paper_name:row.update(paper[paper_name])
        rows.append(row); primary.append(row)
    for label, chosen, columns in [
        ('Mean: 13 RGB / box splits', primary, OUR_COLUMNS[:5]),
        (f'Mean: {len(captions)} matched caption splits',
         [r for r in primary if r.get('Caption') is not None], OUR_COLUMNS[:8])]:
        row = make_row('Aggregate', label, 'mixed', 'equal-weight split mean', f'{len(chosen)} splits',
            notes='Equal weight per named split, including both VLM4D splits. Excludes ScanNet and historical runs. Not the 28-benchmark paper average.',
            sources=[main_root / 'summary.json', *[p / 'summary.json' for p in caption_roots if (p / 'summary.json').exists()]])
        for column in columns:row[column] = statistics.mean(r[column] for r in chosen)
        rows.append(row)
    scan_records = groups['scannet_val88_oracle']
    scan_notes = ('88 ScanNet scenes; GT 2D proposals, categories and poses for Full/Small. '
        'All geometry lists contain 1,224 objects but no explicit tracks or appearance timestamps. '
        'Instance IDs/order differ across GT and predicted lists: not a strict coordinate-only oracle. '
        'Paper VSI scores cover a different scope.')
    scan = fill(make_row('Current benchmark', 'VSI-Bench / ScanNet val88', 'static scene video',
        'overall_score', '2071 questions / arm', notes=scan_notes,
        sources=[r['result'] for r in scan_records] + [f'{PAPER}#page=11']), scan_records)
    scan.update(paper['VSI-Bench']); rows.append(scan)
    # Details already discussed in the conversation: all ScanNet categories and DSI tasks.
    for metric in result(scan_records[0])['statistics']['category_results']:
        if metric == 'overall':continue
        row = make_row('ScanNet QA detail', 'VSI-Bench / ScanNet val88', 'static scene video', metric,
            'category of 2071 questions', notes=scan_notes, sources=[r['result'] for r in scan_records])
        for record in scan_records:row[ARM[(record['variant'], record['mode'])]] = 100 * result(record)['statistics']['category_results'][metric]
        rows.append(row)
    row = make_row('Diagnostic aggregate', 'VSI-Bench / ScanNet val88', 'static scene video',
        'mean excluding appearance order', '7 categories', notes='Diagnostic only; not the official overall metric.',
        sources=[r['result'] for r in scan_records])
    for record in scan_records:
        vals = result(record)['statistics']['category_results']
        row[ARM[(record['variant'], record['mode'])]] = 100 * statistics.mean(
            value for key, value in vals.items() if key not in ['overall', 'obj_appearance_order_accuracy'])
    rows.append(row)
    dsi_records = groups['dsi_bench_all']
    for kind in ['per_type', 'robust_accuracy']:
        for metric, value in result(dsi_records[0])['statistics'][kind].items():
            row = make_row('DSI detail', 'DSI-Bench all augmentations', 'object / camera dynamics',
                metric if kind == 'per_type' else f'robust_accuracy_{metric}',
                f"{value['count']} augmented questions" if kind == 'per_type' else '1769 base items',
                'caption incomplete' if 'dsi_bench_all' not in captions else 'complete',
                notes='Native DSI evaluator; caption columns only filled after complete-split validation.',
                sources=[r['result'] for r in dsi_records])
            for record in dsi_records:
                value = result(record)['statistics'][kind][metric]
                row[ARM[(record['variant'], record['mode'])]] = 100 * (value['accuracy'] if kind == 'per_type' else value)
            rows.append(row)
    # Direct geometric metrics retain their own units, separate from LLM QA scores.
    geometry_root = ROOT / 'results/full_small_20260914_r2'
    geometry_logs = {}; geometry_metrics = {}; categories = {}
    for variant in ['full_epoch64', 'small_epoch34']:
        path = geometry_root / 'runs' / f'{variant}_scannet_val88_gt2d_all' / 'metrics.log'
        text = path.read_text(); geometry_logs[variant] = path
        geometry_metrics[variant] = {k: float(v) for k, v in re.findall(
            r'^\s+([\w@.]+):\s+([0-9.]+)\s*$', text.split('Evaluation Results:')[-1], re.M)}
        macro = text.split('--- Macro F1')[1].split('Evaluation Results:')[0]
        for threshold, score in re.findall(r'F1@([\d.]+):\s+([\d.]+)', macro):
            geometry_metrics[variant][f'macro_f1@{threshold}'] = float(score)
        categories[variant] = {}
        for line in text.splitlines():
            if '|' not in line:continue
            pieces = [x.strip() for x in line.split('|')]
            if len(pieces) != 5 or not re.fullmatch(r'\d+\s+\d+\s+\d+\s+\d+', pieces[1]):continue
            categories[variant][pieces[0]] = {'counts': list(map(int, pieces[1].split())),
                'f1': [float(x.split()[-1]) for x in pieces[2:]]}
        assert len(categories[variant]) == 33
    for metric in geometry_metrics['full_epoch64']:
        percentage = '@' in metric
        counts = metric.startswith(('gt_num', 'pred_num', 'tp', 'fp', 'fn'))
        units = '0–100' if percentage else ('count' if counts else '0–1' if metric in ['iou_3d', 'scale_error'] else 'ScanNet coordinate units')
        row = make_row('Direct 3D geometry', 'ScanNet val88', 'GT 2D proposals + poses', metric,
            '1224 objects / 88 scenes', units=units,
            direction='lower' if metric.endswith('_error') else 'context' if counts else 'higher',
            notes='No LLM. Full/Small columns here are direct geometry measurements. Log precision retained; scale error is 1 minus size-only IoU.',
            sources=geometry_logs.values())
        for variant, column in [('full_epoch64', 'Full boxes'), ('small_epoch34', 'Small boxes')]:
            row[column] = geometry_metrics[variant][metric] * (100 if percentage else 1)
        rows.append(row)
    for category in categories['full_epoch64']:
        for i, threshold in enumerate([0.1, 0.25, 0.5]):
            full = categories['full_epoch64'][category]; small = categories['small_epoch34'][category]
            assert full['counts'][0] == small['counts'][0]
            row = make_row('3D geometry category', f'ScanNet val88 / {category}', 'GT 2D proposals + poses',
                f'F1@{threshold}', f"{full['counts'][0]} GT objects", notes=(
                    f"No LLM; per-category precision = recall = F1 in these runs. TP Full={full['counts'][i+1]}, "
                    f"Small={small['counts'][i+1]}. Values retain 3-decimal log precision before scaling."),
                sources=geometry_logs.values())
            row.update({'Full boxes': 100 * full['f1'][i], 'Small boxes': 100 * small['f1'][i]}); rows.append(row)
    # Earlier complete runs are preserved as historical, never substituted into current scores.
    historical_count = 0
    for folder, label in [('media_models_20260915', 'Historical: initial unconstrained answers'),
                          ('media_native_answers_20260915', 'Historical: native-answer repair')]:
        source = ROOT / 'results' / folder
        state = read_json(source / 'status.json')
        for benchmark, info in state['jobs'].items():
            records = info.get('results', [])
            if not records:continue
            display = next((b[1] for b in BENCHES if b[0] == benchmark), benchmark)
            row = make_row(label, display, 'geometry text', 'native benchmark score',
                notes='Historical protocol; excluded from current means. Do not interpret changes as caused only by truncation.',
                sources=[r['result'] for r in records])
            extra = []
            for record in records:
                d = read_json(record['result']); historical_count += 1
                row['N / unit'] = f"{d['num_samples']} questions / arm"
                row['Full boxes' if record['variant'] == 'full_epoch64' else 'Small boxes'] = d['primary_metric']['score_100']
                invalid = sum(v for k, v in d['status_counts'].items() if k != 'ok')
                extra.append(f"{record['variant']}: max_tokens={d['config']['max_tokens']}, invalid={invalid}")
            row['Protocol / caveats'] += ' ' + '; '.join(extra); rows.append(row)
    original_scan = read_json(geometry_root / 'summary.json')['rows']
    for protocol in sorted({r['metric_protocol'] for r in original_scan}):
        subset = [r for r in original_scan if r['metric_protocol'] == protocol]
        row = make_row('Historical: ScanNet QA', 'VSI-Bench / ScanNet val88', 'GT 2D + predicted 3D + GT pose',
            protocol, '2071 questions / arm', notes='Earlier prompting/output protocol; Full has 2 invalid completions, Small has 1. Not the final matched ablation.',
            sources=[r['result_path'] for r in subset])
        for record in subset:
            row['Full boxes' if record['run_label'].startswith('full') else 'Small boxes'] = record['score_100']
        rows.append(row); historical_count += len(subset)
    for name in ['q_spatial_bench_QSpatial_plus', 'embspatial_bench_test']:
        audit = read_json(ROOT / 'results/media_qa4096_20260915_waiting_archive/baseline_truncation_audit.json')
        subset = [r for r in audit if r['job'] == name]
        row = make_row('Historical diagnostic (not measured)', name, 'geometry text',
            'all originally truncated answers assumed correct', status='counterfactual only',
            notes='Holds nontruncated answers fixed and changes all truncated answers to correct. Not an observed score or a bound on later reruns.',
            sources=[ROOT / 'results/media_qa4096_20260915_waiting_archive/baseline_truncation_audit.json'])
        for r in subset:
            row['Full boxes' if r['variant'] == 'full_epoch64' else 'Small boxes'] = r['baseline_score_100'] + 100 * r['invalid_count'] / r['samples']
            row['N / unit'] = f"{r['samples']} questions"
        rows.append(row)
    # All other benchmark scores visible in the supplied Table 4 remain reference-only rows.
    used = {b[3] for b in BENCHES if b[3]} | {'VSI-Bench'}
    for benchmark in PAPER_BENCHES:
        if benchmark in used:continue
        row = make_row('Paper reference only', benchmark, 'published setting',
            'reported overall average' if benchmark == 'Overall Average' else 'reported benchmark score',
            '28 benchmarks' if benchmark == 'Overall Average' else '', status='not run by us',
            notes='PhysBrain 1.5 Table 4. Closed-source models excluded from the paper ranking. No local score implied.',
            sources=[f'{PAPER}#page=11']); row.update(paper[benchmark]); rows.append(row)
    # Write all representations from the same records; no missing score is zero-filled.
    normalized = [{c: row.get(c) for c in COLUMNS} for row in rows]
    assert len({(r['Section'], r['Benchmark / split'], r['Metric']) for r in normalized}) == len(normalized)
    for row in normalized:
        for col in OUR_COLUMNS + PAPER_MODELS:
            value = row[col]
            assert value is None or (isinstance(value, (int, float)) and math.isfinite(value))
            if value is not None and (row['Units'] == '0–100' or col in PAPER_MODELS):assert 0 <= value <= 100 + 1e-7
    metadata = {'updated_at': timestamp, 'rows': len(rows), 'current_rgb_box_arms': len(summary),
        'current_caption_arms': sum(len(v) for v in captions.values()),
        'current_valid_responses': sum(r['samples'] for r in selected),
        'historical_qa_arms': historical_count, 'paper_benchmarks': 28, 'paper_models': 12,
        'validation': 'Source scores reconciled, completed QA outputs checked for status/finish and sample counts; averages recalculated. Workbook structure checked programmatically; no GUI rendering claimed.',
        'exclusions': 'Smoke tests, pilots, progress estimates and raw per-question predictions are not benchmark scores.',
        'note': 'This is a snapshot. Re-run the script to refresh pending DSI results. All sources are local or the supplied PDF.'}
    write_json(out / 'all_results.json', {'metadata': metadata, 'columns': COLUMNS, 'rows': normalized})
    write_json(out / 'paper_table4.json', paper)
    with (out / 'all_results.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=COLUMNS); writer.writeheader(); writer.writerows(normalized)
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'All results'; ws.append(COLUMNS)
    for row in normalized:ws.append([row[c] for c in COLUMNS])
    table = Table(displayName='AllEvaluationResults', ref=f'A1:{get_column_letter(len(COLUMNS))}{len(rows)+1}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(table)
    ws.freeze_panes = 'I2'; ws.sheet_view.zoomScale = 70
    for cell in ws[1]:
        cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='16324F')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 58
    for i, column in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 18 if column in OUR_COLUMNS + PAPER_MODELS else 27
    ws.column_dimensions['B'].width = 34; ws.column_dimensions['D'].width = 34
    ws.column_dimensions[get_column_letter(len(COLUMNS)-1)].width = 85
    ws.column_dimensions[get_column_letter(len(COLUMNS))].width = 75
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.0000' if row[4].value in ['0–1', 'ScanNet coordinate units'] else '0' if row[4].value == 'count' else '0.00'
        if row[7].value != 'complete':row[7].fill = PatternFill('solid', fgColor='FFF1CC')
        if row[0].value == 'Aggregate':
            for cell in row:cell.font = Font(bold=True)
    ws.print_title_rows = '1:1'; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.save(out / 'all_results.xlsx')
    check = load_workbook(out / 'all_results.xlsx', read_only=True, data_only=True)
    assert check.sheetnames == ['All results'] and check.active.max_row == len(rows)+1
    assert list(next(check.active.values)) == COLUMNS
    for expected, observed in zip(normalized, list(check.active.values)[1:]):
        for key, value in zip(COLUMNS, observed):
            target = expected[key]
            assert math.isclose(target, value, abs_tol=1e-10) if isinstance(target, (int,float)) else (target or None) == value
    check.close()
    notes = [f'# All conversation evaluation results', '', f'Snapshot: {timestamp}.', '',
        f'{len(rows)} rows in one table. Current validated QA: {len(summary)} RGB/box/GT arms and '
        f'{metadata["current_caption_arms"]}/39 caption arms; {metadata["current_valid_responses"]:,} responses.', '',
        'Our QA backbone is Qwen3.5-9B. Full = spatial_encoder_v2 epoch 64; Small = '
        'spatial_encoder_v2_small epoch 34. The backbone is multimodal; box/caption arms send text only. '
        'Direct geometry rows do not use an LLM. Paper columns are reported values, not reruns. '
        'Blank cells mean unavailable or inapplicable, never zero.', '',
        'Filter Section to distinguish current benchmarks, aggregates, ScanNet/DSI details, '
        'direct geometry, historical protocols and paper-only references. QA rates and F1 use 0–100; '
        'IoU, scale error and coordinate errors retain the Units column. No mean mixes these units.', '',
        'The 13-split mean and matched-caption mean have different denominators. Neither is '
        'the paper 28-benchmark average. DSI caption scores remain blank until all three arms complete.', '',
        f'Files: [Excel](all_results.xlsx), [CSV](all_results.csv), [JSON with provenance](all_results.json).', '',
        '| ' + ' | '.join(COLUMNS) + ' |', '| ' + ' | '.join(['---'] * len(COLUMNS)) + ' |']
    notes += ['| ' + ' | '.join(safe_md(r[c]) for c in COLUMNS) + ' |' for r in normalized]
    (out / 'all_results.md').write_text('\n'.join(notes)+'\n')
    preview_columns = ['Benchmark / split', *OUR_COLUMNS, 'RynnBrain1.1 9B (reported)', 'PhysBrain1.5 8B (reported)']
    preview_rows = primary + [scan] + [r for r in rows if r['Section'] == 'Aggregate']
    preview = ['| ' + ' | '.join(preview_columns) + ' |', '| ' + ' | '.join(['---'] + ['---:'] * (len(preview_columns)-1)) + ' |']
    preview += ['| ' + ' | '.join(safe_md(r.get(c)) for c in preview_columns) + ' |' for r in preview_rows]
    (out / 'main_results_preview.md').write_text('\n'.join(preview)+'\n')
    print(metadata, flush=True)
    print(out / 'all_results.xlsx', flush=True)


if __name__ == '__main__':
    main()
